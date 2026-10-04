"""Build the analyst-facing triage workbook."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
TIER_FILL = {"HIGH": "FFC7CE", "MEDIUM": "FFEB9C", "LOW": "C6EFCE"}
TIER_FONT = {"HIGH": "9C0006", "MEDIUM": "9C6500", "LOW": "006100"}

TRIAGE_COLUMNS = [
    "PlayerID", "Shop", "Source", "RiskTier", "FinalScore", "ReviewDecision", "Priority",
    "CriticalSignal", "CriticalCodes", "TopSignal1", "TopSignal2", "TopSignal3",
    "SignalCount", "DecisionPath", "LLM_Typology", "LLM_Summary", "LLM_KeyEvidence",
    "LLM_RecommendedActions", "LLM_Confidence", "LLM_Engine", "DataWarning",
    "deposit7d", "deposit30d", "deposit90d", "withdrawal7d", "withdrawal30d", "withdrawal90d",
    "bet7d", "bet30d", "bet90d", "loss7d", "loss30d", "loss90d",
    "loss_ratio_7d", "loss_ratio_30d", "loss_ratio_90d",
    "bet_dep_30d", "wd_dep_30d", "wd_dep_7d",
    "TicketSystem", "TicketID", "TicketURL", "SLA_Hours", "RunID", "ScoredAt",
]


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(10, longest + 2), max_width)


def write_triage_workbook(triage_rows: list[dict], signal_rows: list[dict],
                          run_meta: dict, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    df = pd.DataFrame(triage_rows)
    for c in TRIAGE_COLUMNS:
        if c not in df.columns:
            df[c] = None
    order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    df = (df.assign(_o=df["RiskTier"].map(order))
            .sort_values(["_o", "FinalScore"], ascending=[True, False])
            .drop(columns="_o"))[TRIAGE_COLUMNS]

    counts = df["RiskTier"].value_counts().to_dict()
    summary = pd.DataFrame([
        {"Metric": "Run ID", "Value": run_meta.get("run_id")},
        {"Metric": "Generated at (UTC)", "Value": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")},
        {"Metric": "Source email", "Value": run_meta.get("source_email", "")},
        {"Metric": "Input files", "Value": ", ".join(run_meta.get("input_files", []))},
        {"Metric": "Framework version", "Value": run_meta.get("framework_version", "")},
        {"Metric": "Explanation engine", "Value": run_meta.get("explain_mode", "")},
        {"Metric": "Players in review population", "Value": len(df)},
        {"Metric": "HIGH risk", "Value": counts.get("HIGH", 0)},
        {"Metric": "MEDIUM risk", "Value": counts.get("MEDIUM", 0)},
        {"Metric": "LOW risk", "Value": counts.get("LOW", 0)},
        {"Metric": "Critical signal cases", "Value": int(df["CriticalSignal"].fillna(False).astype(bool).sum())},
        {"Metric": "Manual review required", "Value": int((df["ReviewDecision"] == "YES").sum())},
        {"Metric": "Tickets created", "Value": run_meta.get("tickets_created", 0)},
        {"Metric": "Records with data warnings", "Value": int(df["DataWarning"].notna().sum())},
    ])

    sig = pd.DataFrame(signal_rows) if signal_rows else pd.DataFrame(
        columns=["PlayerID", "RiskTier", "SignalCode", "Kind", "Family", "RawScore", "Reason", "Evidence"])

    by_signal = (sig.groupby(["SignalCode", "Kind", "Family"], as_index=False)
                    .agg(Players=("PlayerID", "nunique"), AvgRawScore=("RawScore", "mean"))
                    .sort_values("Players", ascending=False)) if not sig.empty else pd.DataFrame()

    with pd.ExcelWriter(out_path, engine="openpyxl") as xl:
        summary.to_excel(xl, sheet_name="Run Summary", index=False)
        df.to_excel(xl, sheet_name="Triage", index=False)
        sig.to_excel(xl, sheet_name="Signal Detail", index=False)
        if not by_signal.empty:
            by_signal.to_excel(xl, sheet_name="Signal Stats", index=False)

        wb = xl.book
        thin = Side(style="thin", color="D9D9D9")
        for name in wb.sheetnames:
            ws = wb[name]
            for cell in ws[1]:
                cell.font = Font(name=FONT, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3864")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name=FONT, size=10)
                    cell.border = Border(bottom=thin)
                    cell.alignment = Alignment(vertical="top", wrap_text=(name == "Triage"))
            ws.freeze_panes = "A2"
            _autosize(ws)

        ws = wb["Triage"]
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["P"].width = 70   # LLM_Summary
        ws.column_dimensions["Q"].width = 60   # LLM_KeyEvidence
        ws.column_dimensions["R"].width = 55   # LLM_RecommendedActions
        tier_col = get_column_letter(TRIAGE_COLUMNS.index("RiskTier") + 1)
        rng = f"{tier_col}2:{tier_col}{ws.max_row}"
        for tier, fill in TIER_FILL.items():
            ws.conditional_formatting.add(rng, CellIsRule(
                operator="equal", formula=[f'"{tier}"'],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(name=FONT, bold=True, color=TIER_FONT[tier])))
        score_col = get_column_letter(TRIAGE_COLUMNS.index("FinalScore") + 1)
        for c in ws[score_col][1:]:
            c.number_format = "0"
        for name in ("loss_ratio_7d", "loss_ratio_30d", "loss_ratio_90d", "bet_dep_30d", "wd_dep_30d", "wd_dep_7d"):
            col = get_column_letter(TRIAGE_COLUMNS.index(name) + 1)
            for c in ws[col][1:]:
                c.number_format = "0.00"
        for name in ("deposit7d", "deposit30d", "deposit90d", "withdrawal7d", "withdrawal30d",
                     "withdrawal90d", "bet7d", "bet30d", "bet90d", "loss7d", "loss30d", "loss90d"):
            col = get_column_letter(TRIAGE_COLUMNS.index(name) + 1)
            for c in ws[col][1:]:
                c.number_format = '#,##0;(#,##0);-'
        if ws.max_row > 1:
            t = Table(displayName="TriageTable", ref=ws.dimensions)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            try:
                ws.add_table(t)
            except ValueError:
                pass

        ws = wb["Run Summary"]
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 70
    return out_path
